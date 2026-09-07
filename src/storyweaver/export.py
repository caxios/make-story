"""Export a finished episode, or a whole story, in the formats readers want.

Markdown and plain text are built here; DOCX is built with `python-docx` if it
is installed, and raises a clear instruction if it is not, rather than making
the whole module unimportable.
"""

from __future__ import annotations

import io
import re
from collections.abc import Sequence

from storyweaver.agents.episode_runner import SCENE_BREAK
from storyweaver.models import CharacterProfile, Episode, WorldLore

# The bracketed header `assemble_episode` puts on an episode. Exports build
# their own headings, so it is stripped rather than repeated.
_HEADER = re.compile(r"^\[Episode\s+\d+[^\]]*\]\s*\n+", re.MULTILINE)


def episode_title(episode: Episode) -> str:
    return f"Episode {episode.episode_number}" + (f": {episode.title}" if episode.title else "")


def episode_body(episode: Episode) -> str:
    """The prose alone, with the assembled header and artefacts removed."""
    body = _HEADER.sub("", episode.final_text, count=1).strip()
    # Fix literal escape sequences and API junk from episodes saved before the
    # writer cleanup was added.
    body = body.replace("\\n", "\n").replace("\\t", "\t")
    body = re.sub(
        r"""(?:extras|additional_kwargs|response_metadata|safety_ratings|usage_metadata)"""
        r"""['"]?\s*[:=]\s*\{[^}]{20,}\}""",
        "",
        body,
        flags=re.DOTALL,
    )
    body = re.sub(r"""['"]?signature['"]?\s*[:=]\s*['"][A-Za-z0-9+/=]{40,}['"]""", "", body)
    body = re.sub(r"\n{3,}", "\n\n", body)
    return body.strip()


# --------------------------------------------------------------------------
# Single episode
# --------------------------------------------------------------------------

def to_text(episode: Episode) -> str:
    """Plain text: a title line and the prose."""
    return f"{episode_title(episode)}\n\n{episode_body(episode)}\n"


def to_markdown(episode: Episode, heading_level: int = 1) -> str:
    """Markdown with a heading and centred scene breaks."""
    body = episode_body(episode).replace(SCENE_BREAK, f"<div align=\"center\">{SCENE_BREAK}</div>")
    return f"{'#' * heading_level} {episode_title(episode)}\n\n{body}\n"


# --------------------------------------------------------------------------
# Whole story
# --------------------------------------------------------------------------

def _slug(text: str) -> str:
    """A GitHub-style anchor, so the table of contents actually links."""
    return re.sub(r"[^a-z0-9\s-]", "", text.lower()).strip().replace(" ", "-")


def table_of_contents(episodes: Sequence[Episode]) -> str:
    lines = ["## Contents", ""]
    for episode in episodes:
        title = episode_title(episode)
        words = len(episode_body(episode).split())
        lines.append(f"{episode.episode_number}. [{title}](#{_slug(title)}) — {words:,} words")
    return "\n".join(lines)


def character_index(characters: Sequence[CharacterProfile]) -> str:
    """An appendix a reader can turn to when a name stops being familiar."""
    if not characters:
        return ""
    lines = ["## Appendix: Characters", ""]
    for character in sorted(characters, key=lambda c: c.name):
        lines.append(f"### {character.name}")
        if character.aliases:
            lines.append(f"*Also known as: {', '.join(character.aliases)}*")
        lines.append("")
        if character.appearance:
            lines.append(character.appearance)
            lines.append("")
        if character.personality_summary:
            lines.append(character.personality_summary)
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def world_glossary(world: WorldLore) -> str:
    """The world's rules, places and factions, for the same reason."""
    lines = [f"## Appendix: {world.title}", ""]
    if world.overview:
        lines += [world.overview, ""]

    if world.rules:
        lines += ["### How this world works", ""]
        for rule in world.rules:
            lines.append(f"- **{rule.category}** — {rule.statement}")
        lines.append("")

    if world.locations:
        lines += ["### Places", ""]
        for location in world.locations:
            lines.append(f"- **{location.name}** — {location.description}")
        lines.append("")

    if world.factions:
        lines += ["### Factions", "", ", ".join(world.factions), ""]

    for key, value in world.additional_lore.items():
        lines += [f"### {key.replace('_', ' ').title()}", "", value, ""]

    return "\n".join(lines).rstrip() + "\n"


def assemble_story(
    title: str,
    episodes: Sequence[Episode],
    characters: Sequence[CharacterProfile] = (),
    world: WorldLore | None = None,
    include_appendices: bool = True,
) -> str:
    """The whole book as one Markdown document, with contents and appendices."""
    ordered = sorted(episodes, key=lambda e: e.episode_number)
    parts = [f"# {title}", ""]

    if not ordered:
        parts.append("*No episodes have been written yet.*")
        return "\n".join(parts) + "\n"

    words = sum(len(episode_body(e).split()) for e in ordered)
    parts += [
        f"*{len(ordered)} episode{'s' if len(ordered) != 1 else ''} · {words:,} words*",
        "",
        table_of_contents(ordered),
        "",
        "---",
        "",
    ]

    for episode in ordered:
        parts += [to_markdown(episode, heading_level=2), "", "---", ""]

    if include_appendices:
        if characters:
            parts += [character_index(characters), ""]
        if world is not None:
            parts += [world_glossary(world), ""]

    return "\n".join(parts).rstrip() + "\n"


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

DOCX_MISSING = (
    "DOCX export needs python-docx. Install it with `pip install python-docx`, "
    "or export Markdown instead."
)


def _require_docx():
    try:
        import docx  # noqa: PLC0415 — optional dependency, imported on demand
    except ImportError as error:  # pragma: no cover - exercised by the message test
        raise RuntimeError(DOCX_MISSING) from error
    return docx


def to_docx(
    episode: Episode,
    title: str | None = None,
) -> bytes:
    """One episode as a .docx, with real paragraph styles rather than plain runs."""
    docx = _require_docx()
    document = docx.Document()

    document.add_heading(title or episode_title(episode), level=1)
    _add_prose(document, episode_body(episode))

    return _to_bytes(document)


def story_to_docx(
    title: str,
    episodes: Sequence[Episode],
    characters: Sequence[CharacterProfile] = (),
    world: WorldLore | None = None,
    include_appendices: bool = True,
) -> bytes:
    """The whole story as a .docx, one heading per chapter."""
    docx = _require_docx()
    ordered = sorted(episodes, key=lambda e: e.episode_number)

    document = docx.Document()
    document.add_heading(title, level=0)

    if ordered:
        document.add_heading("Contents", level=1)
        for episode in ordered:
            document.add_paragraph(episode_title(episode), style="List Number")
        document.add_page_break()

    for index, episode in enumerate(ordered):
        document.add_heading(episode_title(episode), level=1)
        _add_prose(document, episode_body(episode))
        if index < len(ordered) - 1:
            document.add_page_break()

    if include_appendices and (characters or world is not None):
        document.add_page_break()
        if characters:
            document.add_heading("Appendix: Characters", level=1)
            for character in sorted(characters, key=lambda c: c.name):
                document.add_heading(character.name, level=2)
                if character.aliases:
                    document.add_paragraph(f"Also known as: {', '.join(character.aliases)}")
                for text in (character.appearance, character.personality_summary):
                    if text:
                        document.add_paragraph(text)
        if world is not None:
            document.add_heading(f"Appendix: {world.title}", level=1)
            if world.overview:
                document.add_paragraph(world.overview)
            if world.rules:
                document.add_heading("How this world works", level=2)
                for rule in world.rules:
                    document.add_paragraph(
                        f"{rule.category} — {rule.statement}", style="List Bullet"
                    )
            if world.locations:
                document.add_heading("Places", level=2)
                for location in world.locations:
                    document.add_paragraph(
                        f"{location.name} — {location.description}", style="List Bullet"
                    )
            if world.factions:
                document.add_heading("Factions", level=2)
                document.add_paragraph(", ".join(world.factions))

    return _to_bytes(document)


def _add_prose(document, body: str) -> None:
    """Add prose paragraph by paragraph, centring the scene breaks."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for block in body.split("\n\n"):
        text = block.strip()
        if not text:
            continue
        paragraph = document.add_paragraph(text)
        if text == SCENE_BREAK:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _to_bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
